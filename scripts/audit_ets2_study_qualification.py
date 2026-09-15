#!/usr/bin/env python3
"""Structural qualification only: no gene effects or program scores."""
from pathlib import Path
import collections
import hashlib
import json
import math
import re
import sys
import zipfile
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw/ets2-study-qualification'
source = RAW / 'GSE255234_raw_counts_for_rna_seq_submission.xlsx'
preparation = json.loads((ROOT / 'config/ets2-validation-gse255234-preparation.json').read_text())
expected_sha256 = preparation['input_sha256'][preparation['matrix']['file']]
if hashlib.sha256(source.read_bytes()).hexdigest() != expected_sha256:
    raise ValueError('Reacquired workbook differs from the preserved preparation input')
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
ws = wb[preparation['matrix']['sheet']]
rows = ws.iter_rows(values_only=True)
header = next(rows)
ids = []
non_integer = negative = non_finite = missing = bad_type = all_zero = 0
sums = [0] * (len(header)-1)
mins = [float('inf')] * (len(header)-1)
maxs = [0] * (len(header)-1)
column_hashes = [hashlib.sha256() for _ in sums]
zero_values = 0
for row in rows:
    ids.append(row[0])
    zero_row = True
    for i,value in enumerate(row[1:]):
        if value is None:
            missing += 1; zero_row=False; continue
        if not isinstance(value,(int,float)) or isinstance(value,bool):
            bad_type += 1; zero_row=False; continue
        if not math.isfinite(value):
            non_finite += 1; zero_row=False; continue
        negative += value < 0
        non_integer += value != int(value)
        sums[i] += value
        mins[i] = min(mins[i],value); maxs[i] = max(maxs[i],value)
        zero_values += value == 0
        zero_row &= value == 0
        column_hashes[i].update((str(value)+'\n').encode())
    all_zero += zero_row
id_counter = collections.Counter(ids)
versionless_counter = collections.Counter(re.sub(r'\.\d+$','',str(g)) for g in ids)
column_digests = [h.hexdigest() for h in column_hashes]
expected = [s['matrix_column_key'] for s in preparation['samples']]
with zipfile.ZipFile(source) as z:
    metadata = {name:z.read(name).decode() for name in ['docProps/core.xml','docProps/app.xml','xl/workbook.xml'] if name in z.namelist()}
report = {
    'source_path':str(source.relative_to(ROOT)),
    'source_bytes':source.stat().st_size,
    'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
    'inspection_scope':'identifiers, column keys, raw-count structural validity, all-zero features and library totals only; no target effects or scores',
    'python':sys.version,
    'openpyxl_version':openpyxl.__version__,
    'sheets':wb.sheetnames,
    'row_count':len(ids), 'sample_column_count':len(header)-1,
    'header':list(header),
    'unique_identifiers':len(id_counter),
    'duplicate_ids':sum(v>1 for v in id_counter.values()),
    'numeric_version_removed_collisions':sum(v>1 for v in versionless_counter.values()),
    'unversioned_ensembl_gene_id_rows':sum(bool(re.fullmatch(r'ENSG\d{11}',str(g))) for g in ids),
    'summary_row_identifiers':[g for g in ids if str(g).startswith('__')],
    'source_id_set_sorted_sha256':hashlib.sha256(('\n'.join(sorted(ids))+'\n').encode()).hexdigest(),
    'numeric_cells':len(ids)*(len(header)-1),
    'non_integer_values':non_integer, 'negative_values':negative,
    'non_finite_values':non_finite, 'missing_values':missing, 'nonnumeric_values':bad_type,
    'all_zero_features':all_zero, 'zero_cells':zero_values,
    'sample_keys_match_preparation_in_order':list(header[1:])==expected,
    'missing_sample_keys':sorted(set(expected)-set(header[1:])),
    'extra_sample_keys':sorted(set(header[1:])-set(expected)),
    'duplicate_sample_keys':len(header[1:])-len(set(header[1:])),
    'exact_duplicate_count_columns':len(column_digests)-len(set(column_digests)),
    'sample_structural_qc':[{'matrix_column_key':key,'library_total':total,'minimum_value':mn,'maximum_value':mx,'source_column_value_sha256':digest} for key,total,mn,mx,digest in zip(header[1:],sums,mins,maxs,column_digests)],
    'all_library_totals_positive':all(v>0 for v in sums),
    'workbook_package_metadata':metadata,
    'count_universe_provenance_status':'Unresolved: structural validity and retained zero features do not establish full pre-abundance-filter gene universe or reconcile RefSeq versus Ensembl provenance.'
}
wb.close()
out = RAW/'matrix-structural-audit.json'
out.write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k not in ['sample_structural_qc','workbook_package_metadata']},indent=2))
