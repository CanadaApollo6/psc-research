"""Reproduce source-preserving PSC tables; no model inference is performed."""

import csv
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PAPER = 'https://www.nature.com/articles/s41467-024-53602-w'


def expanded_rows(body):
    """Expand JATS table spans so secondary signals retain the correct locus."""
    pending = {}
    for row in body.findall('tr'):
        values = {col: value for col, (value, remaining) in pending.items()}
        pending = {col: (value, remaining - 1) for col, (value, remaining) in pending.items() if remaining > 1}
        column = 0
        for cell in row:
            while column in values:
                column += 1
            text = ''.join(cell.itertext()).strip()
            colspan, rowspan = int(cell.get('colspan', '1')), int(cell.get('rowspan', '1'))
            for offset in range(colspan):
                if column + offset in values:
                    raise ValueError('Overlapping table cells')
                values[column + offset] = text
                if rowspan > 1:
                    pending[column + offset] = (text, rowspan - 1)
            column += colspan
        yield [values[col] for col in range(max(values) + 1)]
    if pending:
        raise ValueError('Unfinished table row span')


def parse_loci(path):
    root = ET.parse(path).getroot()
    body = root.find(".//table-wrap[@id='Tab1']/table/tbody")
    if body is None:
        raise ValueError('Expected Table 1 is missing')
    records = []
    for index, row in enumerate(expanded_rows(body), start=1):
        if len(row) != 9:
            raise ValueError(f'Unexpected Table 1 width at data row {index}')
        chromosome, gene, lead, lead_pos, signal, candidate, candidate_pos, probability, count = row
        if not re.fullmatch(r'rs\d+', lead) or not re.fullmatch(r'rs\d+', candidate):
            raise ValueError(f'Unexpected SNP identifier at data row {index}')
        if not 0 <= float(probability) <= 1 or int(count) < 1:
            raise ValueError(f'Invalid probability or set size at data row {index}')
        records.append(dict(
            source_table_data_row=index, chromosome=chromosome,
            candidate_gene_as_published=gene, lead_gwas_snp=lead,
            lead_position_b37=int(lead_pos), signal=int(signal),
            highest_pp_snp=candidate, highest_pp_position_b37=int(candidate_pos),
            posterior_probability_as_published=probability,
            credible_set_size_as_published=int(count), source_url=PAPER + '#Tab1',
            model_input_status='Not normalized; build and alleles require verification',
        ))
    return records


def parse_colocalisation(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records, context = [], None
    try:
        sheet = book['Supplementary_Data_2']
        expected_header = ('Chr', 'Lead GWAS SNP', 'eGene', 'Molecular QTL Type', 'Tissue', 'Tissue state', 'PP4', 'Risk allele', 'Beta', 'p-value', 'Risk allele', 'Beta', 'p-value')
        if tuple(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))) != expected_header:
            raise ValueError('Supplementary Data 2 headers changed')
        for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
            if not any(value is not None for value in row):
                continue
            if row[0] is not None:
                context = [row[0], row[1], row[7], row[8], row[9]]
            if context is None or len(row) != 13 or row[6] is None:
                raise ValueError(f'Incomplete locus context at spreadsheet row {number}')
            chromosome, lead, risk_allele, beta, pvalue = context
            records.append(dict(
                source_spreadsheet_row=number, chromosome=chromosome, lead_gwas_snp=lead,
                egene_as_published=row[2], molecular_qtl_type=row[3], tissue=row[4],
                tissue_state_as_published=row[5], colocalisation_pp4=row[6],
                gwas_risk_allele=risk_allele, gwas_beta=beta, gwas_pvalue=pvalue,
                functional_risk_allele_as_published=row[10], functional_beta=row[11],
                functional_pvalue_as_published=row[12], source_url=PAPER,
            ))
    finally:
        book.close()
    return records


def parse_geo(text, series):
    records = []
    for block in text.split('^SAMPLE = ')[1:]:
        sample = block.splitlines()[0].strip()
        fields = {}
        for line in block.splitlines():
            if line.startswith('!Sample_') and ' = ' in line:
                key, value = line.split(' = ', 1)
                fields.setdefault(key, []).append(value)
        title = fields['!Sample_title'][0]
        donor = title.split('_')[0]
        cohort = 'PSC' if donor.startswith('PSC') else 'PBC' if donor.startswith('PBC') else 'Control label'
        assay = 'single-nucleus' if '_SN_' in title else 'single-cell' if '_SC_' in title else 'Unclassified'
        chemistry = '5-prime' if '_5pr' in title else '3-prime' if '_3pr' in title else 'Unclassified'
        description = '; '.join(fields.get('!Sample_description', []))
        flags = []
        if chemistry == '5-prime' and '3pr' in description:
            flags.append('Title says 5pr; description says 3pr')
        if assay == 'single-nucleus' and 'scRNA-seq' in description:
            flags.append('Title indicates nuclei; description says scRNA-seq')
        records.append(dict(
            series=series, sample=sample, title=title, donor_label_from_title=donor,
            cohort_from_title=cohort, assay_from_title=assay, chemistry_from_title=chemistry,
            source_name='; '.join(fields.get('!Sample_source_name_ch1', [])),
            description=description, flags='; '.join(flags),
            source_url=f'https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc={sample}',
        ))
    return records


def write_csv(path, records):
    if not records:
        raise ValueError(f'No records for {path.name}')
    with path.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]), lineterminator='\n')
        writer.writeheader()
        writer.writerows(records)


def main():
    manifest = json.loads((ROOT / 'config/sources.json').read_text())
    for source in manifest['sources']:
        path = ROOT / 'data/raw' / source['file']
        if hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
            raise ValueError(f"Unverified source: {source['id']}")
    loci = parse_loci(ROOT / 'data/raw/PMC11541731.xml')
    molecular = parse_colocalisation(ROOT / 'data/raw/41467_2024_53602_MOESM5_ESM.xlsx')
    libraries = []
    for series in ['GSE243977', 'GSE247128']:
        libraries.extend(parse_geo((ROOT / f'data/raw/{series}.soft').read_text(), series))
    counts = Counter(row['cohort_from_title'] for row in libraries)
    audit = {
        'source_snapshot_date': manifest['snapshot_date'],
        'table1_association_signals': len(loci),
        'table1_distinct_loci': len({row['lead_gwas_snp'] for row in loci}),
        'table1_largest_credible_set_as_published': max(row['credible_set_size_as_published'] for row in loci),
        'table1_scope': 'Top candidate per signal only; complete credible-set members are not imported',
        'molecular_colocalisation_rows': len(molecular),
        'geo_library_records': len(libraries), 'geo_libraries_by_title_cohort': dict(counts),
        'geo_unique_donor_labels_from_titles': {cohort: len({row['donor_label_from_title'] for row in libraries if row['cohort_from_title'] == cohort}) for cohort in sorted(counts)},
        'geo_records_with_cross_field_flags': sum(bool(row['flags']) for row in libraries),
        'alphagenome_predictions_run': False, 'expression_matrices_analyzed': False,
        'limitations': 'Source-derived metadata only. Gene names and contradictory values are preserved. See docs/evidence-log.md.',
    }
    derived = ROOT / 'data/derived'
    derived.mkdir(parents=True, exist_ok=True)
    write_csv(derived / 'psc-signal-summaries.csv', loci)
    write_csv(derived / 'psc-molecular-colocalisation.csv', molecular)
    write_csv(derived / 'geo-library-metadata.csv', libraries)
    (ROOT / 'reports').mkdir(exist_ok=True)
    (ROOT / 'reports/data-audit.json').write_text(json.dumps(audit, indent=2) + '\n')
    print(json.dumps(audit, indent=2))


if __name__ == '__main__':
    main()
